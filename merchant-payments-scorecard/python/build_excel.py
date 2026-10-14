"""
Writes the Excel leadership scorecard.

The "Scorecard" sheet has a month drop-down (cell C4). Every number on it is
an Excel formula (INDEX/MATCH, AVERAGE, OFFSET) that looks up the "Monthly
Trend" sheet, so a reader can pick any month and the page updates itself.
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY = "1F3A5F"
LIGHT = "EEF3F9"
GREEN_TXT, RED_TXT = "1E7B34", "B42318"
thin = Side(style="thin", color="D9D9D9")
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)

# column name in monthly_scorecard -> (label, number format, higher is better?)
MONTHLY_COLS = {
    "month": ("Month", "@", None),
    "active_merchants": ("Active merchants", "#,##0", True),
    "new_merchants": ("New merchants", "#,##0", True),
    "closed_merchants": ("Closed merchants", "#,##0", False),
    "approved_volume": ("Approved volume ($)", "$#,##0", True),
    "volume_mom_pct": ("Volume MoM %", '0.0"%"', True),
    "approved_txns": ("Approved transactions", "#,##0", True),
    "avg_ticket": ("Average ticket ($)", "$#,##0.00", True),
    "approval_rate_pct": ("Approval rate %", '0.00"%"', True),
    "net_revenue": ("Net revenue ($)", "$#,##0", True),
    "revenue_mom_pct": ("Revenue MoM %", '0.0"%"', True),
    "revenue_per_merchant": ("Revenue per active merchant ($)", "$#,##0.00", True),
    "chargeback_rate_pct": ("Chargeback rate %", '0.000"%"', False),
    "resolved_tickets": ("Resolved support tickets", "#,##0", False),
    "pct_within_48h_sla": ("Tickets within 48h SLA %", '0.0"%"', True),
    "avg_csat": ("Average CSAT (1-5)", "0.00", True),
}

SCORECARD_KPIS = ["approved_volume", "net_revenue", "active_merchants", "new_merchants",
                  "closed_merchants", "approval_rate_pct", "avg_ticket", "revenue_per_merchant",
                  "chargeback_rate_pct", "pct_within_48h_sla", "avg_csat"]


def _table(ws, df, start_row=1, formats=None, labels=None):
    formats = formats or {}
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(start_row, j, (labels or {}).get(col, col.replace("_", " ").capitalize()))
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, val in enumerate(row, 1):
            c = ws.cell(i, j, None if pd.isna(val) else val)
            c.border = Border(bottom=thin)
            col = df.columns[j - 1]
            if col in formats:
                c.number_format = formats[col]
            if i % 2 == 0:
                c.fill = PatternFill("solid", fgColor=LIGHT)
    ws.row_dimensions[start_row].height = 32
    for j, col in enumerate(df.columns, 1):
        width = max(len(str(ws.cell(start_row, j).value)) * 0.8, df[col].astype(str).str.len().max() + 2, 10)
        ws.column_dimensions[get_column_letter(j)].width = min(width, 32)
    ws.freeze_panes = ws.cell(start_row + 1, 1)


def _pct_fmt(df):
    return {c: '0.0"%"' if not c.startswith("chargeback") else '0.000"%"' for c in df.columns if c.endswith("pct")}


def build_excel(r: dict, path: Path) -> None:
    wb = Workbook()

    # ---- Monthly Trend (data sheet the scorecard reads from)
    sc = r["monthly_scorecard"]
    trend = wb.active
    trend.title = "Monthly Trend"
    _table(trend, sc, formats={k: v[1] for k, v in MONTHLY_COLS.items()},
           labels={k: v[0] for k, v in MONTHLY_COLS.items()})
    n_rows = len(sc) + 1
    col_letter = {c: get_column_letter(i) for i, c in enumerate(sc.columns, 1)}

    # ---- Scorecard (formula-driven)
    ws = wb.create_sheet("Scorecard", 0)
    ws.sheet_view.showGridLines = False
    ws["B2"] = "Merchant Payments Scorecard"
    ws["B2"].font = Font(bold=True, size=16, color=NAVY)
    ws["B3"] = "Simulated data for a fictional Canadian merchant acquirer. Pick a month in C4 and the page updates."
    ws["B3"].font = Font(italic=True, color="666666")
    ws["B4"] = "Month:"
    ws["B4"].font = Font(bold=True)
    ws["C4"] = sc["month"].iloc[-1]
    ws["C4"].fill = PatternFill("solid", fgColor="FFF4CC")
    ws["C4"].font = Font(bold=True)
    dv = DataValidation(type="list", formula1=f"='Monthly Trend'!$A$2:$A${n_rows}", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("C4")

    headers = ["KPI", "Selected month", "Prior month", "Change", "Avg of prior 6 months", "Vs 6-month avg"]
    for j, h in enumerate(headers, 2):
        c = ws.cell(6, j, h)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.row_dimensions[6].height = 30

    month_range = f"'Monthly Trend'!$A$2:$A${n_rows}"
    for i, key in enumerate(SCORECARD_KPIS, 7):
        label, fmt, higher_better = MONTHLY_COLS[key]
        rng = f"'Monthly Trend'!${col_letter[key]}$2:${col_letter[key]}${n_rows}"
        pos = f"MATCH($C$4,{month_range},0)"
        ws.cell(i, 2, label)
        ws.cell(i, 3, f"=INDEX({rng},{pos})").number_format = fmt
        ws.cell(i, 4, f'=IF({pos}>1,INDEX({rng},{pos}-1),"")').number_format = fmt
        ws.cell(i, 5, f'=IF(D{i}="","",IFERROR(C{i}/D{i}-1,""))').number_format = "+0.0%;-0.0%;0.0%"
        ws.cell(i, 6, f'=IF({pos}>6,AVERAGE(OFFSET({rng.split(":")[0]},{pos}-7,0,6,1)),"")').number_format = fmt
        ws.cell(i, 7, f'=IF(F{i}="","",IFERROR(C{i}/F{i}-1,""))').number_format = "+0.0%;-0.0%;0.0%"
        for j in range(2, 8):
            ws.cell(i, j).border = Border(bottom=thin)
        # green when the change is good for the business, red when it's bad
        good, bad = ("greaterThan", "lessThan") if higher_better else ("lessThan", "greaterThan")
        for col in ("E", "G"):
            ws.conditional_formatting.add(f"{col}{i}", CellIsRule(operator=good, formula=["0.0001"] if good == "greaterThan" else ["-0.0001"], font=Font(color=GREEN_TXT, bold=True)))
            ws.conditional_formatting.add(f"{col}{i}", CellIsRule(operator=bad, formula=["-0.0001"] if bad == "lessThan" else ["0.0001"], font=Font(color=RED_TXT, bold=True)))

    note_row = 7 + len(SCORECARD_KPIS) + 1
    ws.cell(note_row, 2, "Green = better for the business, red = worse. Metric definitions are on the 'Definitions' sheet.").font = Font(italic=True, color="666666")
    for col, w in zip("ABCDEFG", [2, 34, 16, 14, 12, 20, 15]):
        ws.column_dimensions[col].width = w

    # ---- Supporting tables
    for name, title in [("onboarding_by_channel", "Channels"),
                        ("activation_speed_vs_retention", "Activation Speed"),
                        ("industry_performance", "Industries"),
                        ("product_performance", "Products"),
                        ("card_type_mix", "Card Mix"),
                        ("service_levels_by_category", "Service Levels"),
                        ("slow_ticket_vs_churn", "Slow Tickets vs Churn"),
                        ("data_quality_checks", "Data Quality")]:
        df = r[name]
        fmts = _pct_fmt(df)
        fmts.update({c: "$#,##0" for c in df.columns if c in ("approved_volume", "net_revenue")})
        fmts.update({c: "$#,##0.00" for c in df.columns if "revenue_per" in c or c == "avg_ticket"})
        _table(wb.create_sheet(title), df, formats=fmts)

    # ---- Definitions
    d = wb.create_sheet("Definitions")
    defs = [
        ("Active merchant", "A merchant with at least one approved transaction in the month."),
        ("Approved volume", "Dollar value of approved card transactions (before refunds)."),
        ("Approval rate", "Approved transactions ÷ attempted transactions."),
        ("Net revenue", "Processing margin (rate by pricing plan and card type) + monthly platform fee per active merchant."),
        ("Chargeback rate", "Chargebacks ÷ approved credit-card transactions (debit excluded)."),
        ("Within 48h SLA", "Share of resolved tickets closed within 48 hours of opening. Tickets with invalid timestamps are excluded."),
        ("CSAT", "Average 1-5 survey score from merchants who answered the post-ticket survey."),
        ("Activated within 30 days", "Approved merchants whose first transaction came within 30 days of approval."),
        ("Retained 6 months", "Activated merchants whose account was still open 6 months after their first transaction."),
    ]
    _table(d, pd.DataFrame(defs, columns=["metric", "definition"]))
    d.column_dimensions["A"].width = 26
    d.column_dimensions["B"].width = 100

    path.parent.mkdir(exist_ok=True)
    wb.save(path)
